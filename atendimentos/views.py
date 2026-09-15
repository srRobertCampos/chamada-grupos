import calendar
import os
import smtplib
import ssl
from datetime import date
from email.message import EmailMessage
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Configuracao, Grupo, Participacao, Registro, Usuario


def data_selecionada(request):
    try:
        return date.fromisoformat(request.GET.get('data') or request.POST.get('data'))
    except (TypeError, ValueError):
        return timezone.localdate()


@login_required
def hoje(request):
    data = data_selecionada(request)
    grupos = Grupo.objects.filter(ativo=True, dia_semana=data.weekday()).select_related('profissional_1', 'profissional_2')
    cards = []
    for grupo in grupos:
        participantes = Participacao.objects.filter(grupo=grupo, ativo=True, usuario__ativo=True).select_related('usuario').order_by('usuario__nome_completo')
        registros = {r.usuario_id: r for r in Registro.objects.filter(grupo=grupo, data=data)}
        cards.append((grupo, [(p.usuario, registros.get(p.usuario_id)) for p in participantes]))
    return render(request, 'atendimentos/hoje.html', {'data': data, 'cards': cards, 'tipos': Registro.TIPOS})


@login_required
@require_POST
def registrar(request, grupo_id):
    data = data_selecionada(request)
    grupo = get_object_or_404(Grupo, pk=grupo_id, ativo=True, dia_semana=data.weekday())
    ids = set(Participacao.objects.filter(grupo=grupo, ativo=True, usuario__ativo=True).values_list('usuario_id', flat=True))
    validos = dict(Registro.TIPOS)
    for usuario_id in ids:
        tipo = request.POST.get(f'tipo_{usuario_id}', '')
        if tipo in validos:
            Registro.objects.update_or_create(data=data, grupo=grupo, usuario_id=usuario_id, defaults={'tipo': tipo})
        elif not tipo:
            Registro.objects.filter(data=data, grupo=grupo, usuario_id=usuario_id).delete()
    messages.success(request, f'Chamada de {grupo.nome} salva. A folha mensal foi atualizada.')
    if Configuracao.objects.filter(email_destino__gt='').exists() and os.environ.get('SMTP_HOST') and os.environ.get('SMTP_FROM'):
        try:
            enviar_planilha(data.year, data.month)
        except (OSError, smtplib.SMTPException) as erro:
            messages.error(request, f'Chamada salva, mas o envio falhou: {erro}')
        else:
            messages.success(request, 'Folha mensal enviada por e-mail.')
    return redirect(f'/?data={data.isoformat()}')


@login_required
def extra(request):
    data = data_selecionada(request)
    if request.method == 'POST':
        usuario = get_object_or_404(Usuario, pk=request.POST.get('usuario'), ativo=True)
        grupo = get_object_or_404(Grupo, pk=request.POST.get('grupo'), ativo=True)
        tipo = request.POST.get('tipo')
        if tipo not in (Registro.FAMILIAR, Registro.EXTRA):
            messages.error(request, 'Tipo de atendimento inválido.')
        else:
            Registro.objects.update_or_create(data=data, grupo=grupo, usuario=usuario, defaults={'tipo': tipo, 'observacao': request.POST.get('observacao', '')[:250]})
            messages.success(request, 'Atendimento registrado e incluído na folha mensal.')
            return redirect(f'/?data={data.isoformat()}')
    return render(request, 'atendimentos/extra.html', {'data': data, 'usuarios': Usuario.objects.filter(ativo=True), 'grupos': Grupo.objects.filter(ativo=True)})


def dados_folha(ano, mes):
    registros = Registro.objects.filter(data__year=ano, data__month=mes).select_related('usuario', 'grupo').order_by('grupo__horario', 'grupo__nome', 'usuario__nome_completo')
    linhas = {}
    for participacao in Participacao.objects.filter(ativo=True, grupo__ativo=True, usuario__ativo=True).select_related('grupo', 'usuario'):
        linhas[(participacao.grupo_id, participacao.usuario_id)] = {'grupo': str(participacao.grupo), 'usuario': participacao.usuario.nome_completo, 'dias': {}}
    for registro in registros:
        chave = (registro.grupo_id, registro.usuario_id)
        if chave not in linhas:
            linhas[chave] = {'grupo': str(registro.grupo), 'usuario': registro.usuario.nome_completo, 'dias': {}}
        linhas[chave]['dias'][registro.data.day] = registro.tipo
    return list(linhas.values()), calendar.monthrange(ano, mes)[1]


def periodo(request):
    hoje = timezone.localdate()
    try:
        ano = int(request.GET.get('ano', hoje.year))
        mes = int(request.GET.get('mes', hoje.month))
        if not 1 <= mes <= 12 or not 2000 <= ano <= 2100:
            raise ValueError
        return ano, mes
    except ValueError:
        return hoje.year, hoje.month


@login_required
def folha(request):
    ano, mes = periodo(request)
    linhas, dias = dados_folha(ano, mes)
    return render(request, 'atendimentos/folha.html', {'ano': ano, 'mes': mes, 'linhas': linhas, 'dias': range(1, dias + 1)})


def arquivo_excel(ano, mes):
    linhas, dias = dados_folha(ano, mes)
    livro = Workbook()
    aba = livro.active
    aba.title = f'{mes:02d}-{ano}'
    aba.append([f'Folha mensal — {mes:02d}/{ano}'])
    aba.append(['Grupo', 'Usuário'] + list(range(1, dias + 1)))
    for linha in linhas:
        aba.append([linha['grupo'], linha['usuario']] + [linha['dias'].get(dia, '') for dia in range(1, dias + 1)])
    aba.freeze_panes = 'C3'
    aba.column_dimensions['A'].width = 35
    aba.column_dimensions['B'].width = 32
    for celula in aba[2]:
        celula.fill = PatternFill('solid', fgColor='174A65')
        celula.font = Font(color='FFFFFF', bold=True)
        celula.alignment = Alignment(horizontal='center')
    for coluna in range(3, dias + 3):
        aba.column_dimensions[aba.cell(2, coluna).column_letter].width = 5
    memoria = BytesIO()
    livro.save(memoria)
    return memoria.getvalue()


@login_required
def excel(request):
    ano, mes = periodo(request)
    resposta = HttpResponse(arquivo_excel(ano, mes), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resposta['Content-Disposition'] = f'attachment; filename="folha-{ano}-{mes:02d}.xlsx"'
    return resposta


@login_required
@require_POST
def enviar(request):
    ano, mes = periodo(request)
    config = Configuracao.objects.first()
    destino = config.email_destino if config else ''
    if not destino:
        messages.error(request, 'Cadastre o e-mail de destino em Configurações.')
        return redirect(f'/folha/?ano={ano}&mes={mes}')
    host = os.environ.get('SMTP_HOST')
    remetente = os.environ.get('SMTP_FROM')
    if not host or not remetente:
        messages.error(request, 'Configure SMTP_HOST e SMTP_FROM no ambiente antes de enviar.')
        return redirect(f'/folha/?ano={ano}&mes={mes}')
    try:
        enviar_planilha(ano, mes)
    except (OSError, smtplib.SMTPException) as erro:
        messages.error(request, f'Falha no envio: {erro}')
    else:
        messages.success(request, f'Planilha enviada para {destino}.')
    return redirect(f'/folha/?ano={ano}&mes={mes}')


def enviar_planilha(ano, mes):
    destino = Configuracao.objects.first().email_destino
    host = os.environ['SMTP_HOST']
    remetente = os.environ['SMTP_FROM']
    mensagem = EmailMessage()
    mensagem['Subject'] = f'Folha mensal de atendimentos — {mes:02d}/{ano}'
    mensagem['From'] = remetente
    mensagem['To'] = destino
    mensagem.set_content('Segue a folha mensal de atendimentos em anexo.')
    mensagem.add_attachment(arquivo_excel(ano, mes), maintype='application', subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet', filename=f'folha-{ano}-{mes:02d}.xlsx')
    with smtplib.SMTP(host, int(os.environ.get('SMTP_PORT', '587')), timeout=20) as servidor:
        servidor.starttls(context=ssl.create_default_context())
        servidor.login(os.environ.get('SMTP_USER', ''), os.environ.get('SMTP_PASSWORD', ''))
        servidor.send_message(mensagem)
