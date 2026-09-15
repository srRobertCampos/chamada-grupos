from datetime import date, time
from io import BytesIO

from django.contrib.auth.models import User as Conta
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .models import Grupo, Participacao, Profissional, Registro, Usuario
from .views import arquivo_excel


class ChamadaTests(TestCase):
    def setUp(self):
        self.conta = Conta.objects.create_user('profissional', password='senha-forte-123')
        self.client.force_login(self.conta)
        p1 = Profissional.objects.create(nome='Ana')
        p2 = Profissional.objects.create(nome='Bia')
        self.grupo = Grupo.objects.create(nome='Grupo A', dia_semana=0, horario=time(9), profissional_1=p1, profissional_2=p2)
        self.usuario = Usuario.objects.create(nome_completo='João Silva')
        Participacao.objects.create(grupo=self.grupo, usuario=self.usuario)

    def test_chamada_alimenta_planilha_mensal(self):
        data = date(2026, 9, 14)
        resposta = self.client.post(reverse('registrar', args=[self.grupo.id]), {'data': data.isoformat(), f'tipo_{self.usuario.id}': 'FJ'})
        self.assertEqual(resposta.status_code, 302)
        self.assertEqual(Registro.objects.get().tipo, 'FJ')
        aba = load_workbook(BytesIO(arquivo_excel(2026, 9))).active
        self.assertEqual(aba.cell(3, 2).value, 'João Silva')
        self.assertEqual(aba.cell(3, 16).value, 'FJ')

    def test_outro_grupo_nao_pode_ser_registrado_no_dia(self):
        resposta = self.client.post(reverse('registrar', args=[self.grupo.id]), {'data': '2026-09-15', f'tipo_{self.usuario.id}': 'P'})
        self.assertEqual(resposta.status_code, 404)

    def test_folha_inclui_usuarios_sem_lancamento(self):
        aba = load_workbook(BytesIO(arquivo_excel(2026, 9))).active
        self.assertEqual(aba.cell(3, 2).value, 'João Silva')
